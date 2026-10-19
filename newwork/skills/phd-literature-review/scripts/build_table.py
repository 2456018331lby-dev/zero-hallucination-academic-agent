#!/usr/bin/env python3
"""Render a literature comparison table (Markdown + XLSX) from a JSON file.

Input JSON shape:
{
  "meta": {
    "direction": "...", "subtitle": "...",
    "source_note": "...", "rules": "...", "generated": "YYYY-MM-DD",
    "notes": ["extra 取材说明 lines", "..."]        # optional
  },
  "papers": [ {15 fields...}, ... ]
}

Usage:
    python build_table.py --json <literature_table.json> --out-dir <DIR> [--libs <SITE_DIR>]
"""
from __future__ import annotations

import argparse
import json
import os
import sys

FIELDS = [
    ("id", "序号"),
    ("title", "论文题目"),
    ("authors", "作者"),
    ("year", "年份"),
    ("venue", "会议/期刊"),
    ("problem", "研究问题"),
    ("method", "核心方法"),
    ("innovation", "主要创新点"),
    ("setup", "数据集/实验设置"),
    ("results", "关键实验结果"),
    ("pros", "方法优点"),
    ("cons", "方法局限"),
    ("relation", "与本研究方向的关系"),
    ("recommend", "是否建议重点汇报"),
    ("why", "推荐汇报理由"),
]


def cell(s) -> str:
    return str(s).replace("|", "\\|").replace("\n", " ").strip()


def build_md(data: dict) -> str:
    m, ps = data["meta"], data["papers"]
    n = len(ps)
    L = [
        f"# 文献阅读对比表格 · {m['direction']}",
        "",
        f"**子方向**：{m['subtitle']}",
        "",
        f"> **数据来源**：{m['source_note']}",
        "> ",
        f"> **取材原则**：{m['rules']}",
        "",
        f"共 **{n}** 篇。生成日期：{m['generated']}",
        "",
        "## 表 1 · 书目信息总览",
        "",
        "| 序号 | 论文题目 | 作者 | 年份 | 会议/期刊 | 是否建议重点汇报 |",
        "| :---: | :--- | :--- | :---: | :--- | :--- |",
    ]
    for p in ps:
        L.append(f"| {p['id']} | {cell(p['title'])} | {cell(p['authors'])} | {p['year']} | "
                 f"{cell(p['venue'])} | {cell(p['recommend'])} |")
    L += ["", "## 表 2 · 研究内容对比", "",
          "| 序号 | 研究问题 | 核心方法 | 主要创新点 |",
          "| :---: | :--- | :--- | :--- |"]
    for p in ps:
        L.append(f"| {p['id']} | {cell(p['problem'])} | {cell(p['method'])} | {cell(p['innovation'])} |")
    L += ["", "## 表 3 · 实验设置与关键结果", "",
          "| 序号 | 数据集/实验设置 | 关键实验结果 |",
          "| :---: | :--- | :--- |"]
    for p in ps:
        L.append(f"| {p['id']} | {cell(p['setup'])} | {cell(p['results'])} |")
    L += ["", "## 表 4 · 评价、定位与汇报建议", "",
          "| 序号 | 方法优点 | 方法局限 | 与本研究方向的关系 | 是否建议重点汇报 | 推荐汇报理由 |",
          "| :---: | :--- | :--- | :--- | :--- | :--- |"]
    for p in ps:
        L.append(f"| {p['id']} | {cell(p['pros'])} | {cell(p['cons'])} | {cell(p['relation'])} | "
                 f"{cell(p['recommend'])} | {cell(p['why'])} |")
    L += ["", "---", "", "## 附录 · 逐篇完整 15 字段卡片", ""]
    for p in ps:
        L.append(f"### [{p['id']}] {p['title']}")
        L.append("")
        if p.get("file"):
            L.append(f"- **本地文件**：`{p['file']}`")
        for k, label in FIELDS[1:]:
            v = cell(p[k])
            if k == "venue" and p.get("venue_note"):
                v += f"（{cell(p['venue_note'])}）"
            L.append(f"- **{label}**：{v}")
        L.append("")
    return "\n".join(L)


def build_xlsx(data: dict, out_path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    m, ps = data["meta"], data["papers"]
    wb = Workbook()
    ws = wb.active
    ws.title = "文献阅读对比表"

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="D9E2F3")
    alt_fill = PatternFill("solid", fgColor="F2F5FB")

    for c, (_k, label) in enumerate(FIELDS, 1):
        o = ws.cell(row=1, column=c, value=label)
        o.font = Font(bold=True, size=10)
        o.fill = head_fill
        o.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        o.border = border

    for r, p in enumerate(ps, 2):
        for c, (k, _label) in enumerate(FIELDS, 1):
            v = p[k]
            if k == "venue" and p.get("venue_note"):
                v = f"{v}\n（{p['venue_note']}）"
            o = ws.cell(row=r, column=c, value=v)
            o.alignment = Alignment(vertical="top", wrap_text=True)
            o.font = Font(size=9)
            o.border = border
            if r % 2 == 0:
                o.fill = alt_fill

    widths = {"id": 6, "title": 42, "authors": 30, "year": 7, "venue": 26, "problem": 46,
              "method": 56, "innovation": 46, "setup": 52, "results": 46, "pros": 40,
              "cons": 40, "relation": 30, "recommend": 16, "why": 40}
    for c, (k, _label) in enumerate(FIELDS, 1):
        ws.column_dimensions[get_column_letter(c)].width = widths[k]
    ws.row_dimensions[1].height = 30
    for r in range(2, len(ps) + 2):
        ws.row_dimensions[r].height = 130
    ws.freeze_panes = "C2"

    ws2 = wb.create_sheet("取材说明")
    notes = [
        ["项目", "内容"],
        ["研究方向", m["direction"]],
        ["子方向", m["subtitle"]],
        ["数据来源", m["source_note"]],
        ["取材原则", m["rules"]],
        ["生成日期", m["generated"]],
        ["", ""],
    ]
    for i, extra in enumerate(m.get("notes", []), 1):
        notes.append([f"重要说明 {i}", extra])
    notes.append(["重要说明", "所有「文献未说明」标注均表示原文确实未给出该项内容，非遗漏。"])
    for r, row in enumerate(notes, 1):
        for c, v in enumerate(row, 1):
            o = ws2.cell(row=r, column=c, value=v)
            o.alignment = Alignment(vertical="top", wrap_text=True)
            if r == 1:
                o.font = Font(bold=True)
                o.fill = head_fill
    ws2.column_dimensions["A"].width = 16
    ws2.column_dimensions["B"].width = 100

    wb.save(out_path)
    print("wrote", out_path)


def main() -> int:
    ap = argparse.ArgumentParser(description="Render literature table to MD + XLSX.")
    ap.add_argument("--json", required=True, help="literature_table.json")
    ap.add_argument("--out-dir", required=True, help="output directory")
    ap.add_argument("--libs", default=None, help="optional extra sys.path for openpyxl")
    ap.add_argument("--name", default="文献阅读对比表格", help="output basename")
    args = ap.parse_args()

    if args.libs:
        sys.path.insert(0, args.libs)
    os.makedirs(args.out_dir, exist_ok=True)
    with open(args.json, encoding="utf-8") as fh:
        data = json.load(fh)

    md_path = os.path.join(args.out_dir, args.name + ".md")
    with open(md_path, "w", encoding="utf-8") as fh:
        fh.write(build_md(data))
    print("wrote", md_path, os.path.getsize(md_path), "bytes")
    try:
        build_xlsx(data, os.path.join(args.out_dir, args.name + ".xlsx"))
    except ImportError:
        print("openpyxl not importable -> skipped XLSX. "
              "Install isolated: pip install --target <dir> openpyxl, then --libs <dir>",
              file=sys.stderr)
        return 0
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
