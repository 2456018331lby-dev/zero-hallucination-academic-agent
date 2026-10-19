#!/usr/bin/env python3
"""Emit the literature comparison table as Markdown + XLSX from literature_table.json."""
import json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
NW = os.path.normpath(os.path.join(HERE, "..", "..", ".."))
_LIBDIR = os.path.join(NW, "libs", "site-packages")
if os.path.isdir(_LIBDIR):  # optional vendored deps; fall back to installed packages
    sys.path.insert(0, _LIBDIR)

FIELDS = [
    ("id", "序号"),
    ("title", "论文题目"),
    ("authors", "作者"),
    ("year", "年份"),
    ("venue", "会议/期刊"),
    ("problem", "研究问题"),
    ("method", "核心方法"),
    ("innovation", "主要创新点"),
    ("setup", "数据集/实验设置"),
    ("results", "关键实验结果"),
    ("pros", "方法优点"),
    ("cons", "方法局限"),
    ("relation", "与本研究方向的关系"),
    ("recommend", "是否建议重点汇报"),
    ("why", "推荐汇报理由"),
]


def cell(s) -> str:
    return str(s).replace("|", "\\|").replace("\n", " ").strip()


def build_md(data: dict) -> str:
    m, ps = data["meta"], data["papers"]
    L = []
    L.append(f"# 文献阅读对比表格 · {m['direction']}")
    L.append("")
    L.append(f"**子方向**：{m['subtitle']}")
    L.append("")
    L.append(f"> **数据来源**：{m['source_note']}")
    L.append("> ")
    L.append(f"> **取材原则**：{m['rules']}")
    L.append("")
    L.append(f"共 **{len(ps)}** 篇。生成日期：{m['generated']}")
    L.append("")

    L.append("## 表 1 · 书目信息总览")
    L.append("")
    L.append("| 序号 | 论文题目 | 作者 | 年份 | 会议/期刊 | 是否建议重点汇报 |")
    L.append("| :---: | :--- | :--- | :---: | :--- | :--- |")
    for p in ps:
        L.append(f"| {p['id']} | {cell(p['title'])} | {cell(p['authors'])} | {p['year']} | "
                 f"{cell(p['venue'])} | {cell(p['recommend'])} |")
    L.append("")
    L.append("**关于「会议/期刊」一栏的说明**：本地 11 篇 PDF 全文均未标注正式期刊或会议发表信息，"
             "全部为 arXiv 预印本。表中仅列出其 arXiv 编号，不做任何期刊归属的推断。")
    L.append("")

    L.append("## 表 2 · 研究内容对比")
    L.append("")
    L.append("| 序号 | 研究问题 | 核心方法 | 主要创新点 |")
    L.append("| :---: | :--- | :--- | :--- |")
    for p in ps:
        L.append(f"| {p['id']} | {cell(p['problem'])} | {cell(p['method'])} | {cell(p['innovation'])} |")
    L.append("")

    L.append("## 表 3 · 实验设置与关键结果")
    L.append("")
    L.append("| 序号 | 数据集/实验设置 | 关键实验结果 |")
    L.append("| :---: | :--- | :--- |")
    for p in ps:
        L.append(f"| {p['id']} | {cell(p['setup'])} | {cell(p['results'])} |")
    L.append("")

    L.append("## 表 4 · 评价、定位与汇报建议")
    L.append("")
    L.append("| 序号 | 方法优点 | 方法局限 | 与本研究方向的关系 | 是否建议重点汇报 | 推荐汇报理由 |")
    L.append("| :---: | :--- | :--- | :--- | :--- | :--- |")
    for p in ps:
        L.append(f"| {p['id']} | {cell(p['pros'])} | {cell(p['cons'])} | {cell(p['relation'])} | "
                 f"{cell(p['recommend'])} | {cell(p['why'])} |")
    L.append("")

    L.append("---")
    L.append("")
    L.append("## 附录 · 逐篇完整 15 字段卡片")
    L.append("")
    for p in ps:
        L.append(f"### [{p['id']}] {p['title']}")
        L.append("")
        L.append(f"- **本地文件**：`literature/{p['file']}`")
        for k, label in FIELDS[1:]:
            extra = p.get("venue_note") if k == "venue" else None
            v = cell(p[k])
            if extra:
                v += f"（{cell(extra)}）"
            L.append(f"- **{label}**：{v}")
        L.append("")
    return "\n".join(L)


def build_xlsx(data: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "文献阅读对比表"

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="D9E2F3")
    alt_fill = PatternFill("solid", fgColor="F2F5FB")

    for c, (_k, label) in enumerate(FIELDS, 1):
        cellobj = ws.cell(row=1, column=c, value=label)
        cellobj.font = Font(bold=True, size=10)
        cellobj.fill = head_fill
        cellobj.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cellobj.border = border

    for r, p in enumerate(data["papers"], 2):
        for c, (k, _label) in enumerate(FIELDS, 1):
            v = p[k]
            if k == "venue" and p.get("venue_note"):
                v = f"{v}\n（{p['venue_note']}）"
            cellobj = ws.cell(row=r, column=c, value=v)
            cellobj.alignment = Alignment(vertical="top", wrap_text=True)
            cellobj.font = Font(size=9)
            cellobj.border = border
            if r % 2 == 0:
                cellobj.fill = alt_fill

    widths = {"id": 6, "title": 42, "authors": 30, "year": 7, "venue": 26, "problem": 46,
              "method": 56, "innovation": 46, "setup": 52, "results": 46, "pros": 40,
              "cons": 40, "relation": 30, "recommend": 16, "why": 40}
    for c, (k, _label) in enumerate(FIELDS, 1):
        ws.column_dimensions[get_column_letter(c)].width = widths[k]
    ws.row_dimensions[1].height = 30
    for r in range(2, len(data["papers"]) + 2):
        ws.row_dimensions[r].height = 130
    ws.freeze_panes = "C2"

    ws2 = wb.create_sheet("取材说明")
    notes = [
        ["项目", "内容"],
        ["研究方向", data["meta"]["direction"]],
        ["子方向", data["meta"]["subtitle"]],
        ["数据来源", data["meta"]["source_note"]],
        ["取材原则", data["meta"]["rules"]],
        ["生成日期", data["meta"]["generated"]],
        ["", ""],
        ["重要说明 1", "本地 11 篇 PDF 全文均未标注正式期刊/会议发表信息，全部为 arXiv 预印本；表中不做期刊归属推断。"],
        ["重要说明 2", "精读池原有 12 篇，其中 1 篇（Integrated Super-resolution Sensing and Symbiotic "
                      "Communication with 3D Sparse MIMO for Low-Altitude UAV Swarm）因付费墙未下载到本地，"
                      "故未纳入本表——本表只基于本地可获取的 PDF。"],
        ["重要说明 3", "所有「文献未说明」标注均表示原文确实未给出该项内容，非遗漏。"],
    ]
    for r, row in enumerate(notes, 1):
        for c, v in enumerate(row, 1):
            cellobj = ws2.cell(row=r, column=c, value=v)
            cellobj.alignment = Alignment(vertical="top", wrap_text=True)
            if r == 1:
                cellobj.font = Font(bold=True)
                cellobj.fill = head_fill
    ws2.column_dimensions["A"].width = 16
    ws2.column_dimensions["B"].width = 100

    out = os.path.join(HERE, "文献阅读对比表格.xlsx")
    wb.save(out)
    print("wrote", out)


def main() -> None:
    with open(os.path.join(HERE, "literature_table.json"), encoding="utf-8") as fh:
        data = json.load(fh)
    md_path = os.path.join(HERE, "文献阅读对比表格.md")
    with open(md_path, "w", encoding="utf-8") as fh:
        fh.write(build_md(data))
    print("wrote", md_path, os.path.getsize(md_path), "bytes")
    build_xlsx(data)


main()
